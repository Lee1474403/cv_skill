#!/usr/bin/env python3
"""Build an editable DOCX and searchable PDF from verified resume JSON."""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Flowable

from verify_resume import verify


def publish_output(source: Path, target: Path, overwrite: bool = False) -> None:
    """Publish new bytes under the output folder's ACL, not the temp folder's ACL.

    On Windows Python temp directories may have a protected owner-only DACL.
    A same-volume rename preserves that DACL, making files unreadable outside
    the generating account. Create a fresh sibling with normal inheritance.
    """
    staging = target.parent / ("." + target.name + "." + uuid.uuid4().hex + ".publish")
    try:
        with source.open("rb") as reader, staging.open("xb") as writer:
            shutil.copyfileobj(reader, writer)
            writer.flush()
            os.fsync(writer.fileno())
        if overwrite:
            os.replace(staging, target)
        else:
            try:
                # Linking a fresh sibling is atomic and refuses an existing path.
                os.link(staging, target)
            except FileExistsError:
                raise
            except OSError:
                # Shares without hard-link support still get exclusive creation.
                created = False
                try:
                    with staging.open("rb") as reader, target.open("xb") as writer:
                        created = True
                        shutil.copyfileobj(reader, writer)
                except Exception:
                    if created:
                        target.unlink(missing_ok=True)
                    raise
    finally:
        staging.unlink(missing_ok=True)


KINDS = {"education", "experience", "projects", "skills", "awards", "campus",
         "certifications", "publications", "other"}
URL = re.compile(r'https?://[^\s<>\u3002\uff0c\uff1b\uff09\u3001|]+')
DATE = re.compile(
    r"^(\d{4})(?:\.(\d{2}))?(?: - (?:(\d{4})(?:\.(\d{2}))?|(至今|Present)))?"
    r"(?:（预计）| \(Expected\))?$"
)
PLACEHOLDER = re.compile(r"\[(?:待补充|待确认|公司名称|学校名称|姓名)\]|\b(?:TODO|TBD)\b")
REFERENCE_STYLE = json.loads((Path(__file__).resolve().parents[1] / "assets/reference-blue.json").read_text(encoding="utf-8"))


def check_fields(value: object, allowed: set[str], where: str) -> dict:
    if not isinstance(value, dict):
        raise ValueError(f"{where} must be an object.")
    unknown = set(value) - allowed
    if unknown:
        raise ValueError(f"{where}: unknown fields: {', '.join(sorted(unknown))}")
    return value


def check_text(value: object, where: str, required: bool = False) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{where} must be a string.")
    if any(ord(c) < 32 or ord(c) == 127 for c in value):
        raise ValueError(f"{where}: use separate items instead of control characters/newlines.")
    text = value.strip()
    if required and not text:
        raise ValueError(f"{where} cannot be empty.")
    if PLACEHOLDER.search(text):
        raise ValueError(f"{where}: unresolved placeholder; remove or confirm it before export.")
    return text


def check_list(value: object, where: str) -> list[str]:
    if not isinstance(value, list):
        raise ValueError(f"{where} must be a list.")
    return [check_text(v, f"{where}[{i}]", required=True) for i, v in enumerate(value)]


def check_date(value: str, where: str) -> None:
    if not value:
        return
    match = DATE.fullmatch(value)
    if not match:
        raise ValueError(f"{where}: use YYYY[.MM] - YYYY[.MM]/至今/Present, optionally 预计/Expected.")
    start_y, start_m, end_y, end_m, _ = match.groups()
    if start_y == "0000" or end_y == "0000":
        raise ValueError(f"{where}: year zero is invalid.")
    for month in (start_m, end_m):
        if month and not 1 <= int(month) <= 12:
            raise ValueError(f"{where}: invalid month.")
    if end_y:
        if bool(start_m) != bool(end_m):
            raise ValueError(f"{where}: start/end date precision must match.")
        if (int(end_y), int(end_m or 1)) < (int(start_y), int(start_m or 1)):
            raise ValueError(f"{where}: end date precedes start date.")


def validate_profile(raw: object) -> dict:
    data = check_fields(raw, {"name", "language", "target_role", "contacts", "summary", "sections"}, "profile")
    result = {"name": check_text(data.get("name", ""), "name", True)}
    language = data.get("language", "zh")
    if language not in ("zh", "en"):
        raise ValueError("language must be zh or en.")
    result["language"] = language
    result["target_role"] = check_text(data.get("target_role", ""), "target_role")
    for key in ("contacts", "summary"):
        result[key] = check_list(data.get(key, []), key)
    sections = data.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ValueError("sections must be a non-empty list.")
    result["sections"] = []
    for i, raw_section in enumerate(sections):
        where = f"sections[{i}]"
        section = check_fields(raw_section, {"kind", "title", "entries"}, where)
        kind = check_text(section.get("kind", ""), where + ".kind", True)
        if kind not in KINDS:
            raise ValueError(f"{where}: unsupported section kind: {kind}")
        entries = section.get("entries")
        if not isinstance(entries, list) or not entries:
            raise ValueError(f"{where}.entries must be a non-empty list; omit empty sections.")
        clean = {"kind": kind, "title": check_text(section.get("title", ""), where + ".title", True), "entries": []}
        for j, raw_entry in enumerate(entries):
            entry_where = f"{where}.entries[{j}]"
            entry = check_fields(raw_entry, {"heading", "detail", "date", "location", "bullets", "paragraphs", "emphasis"}, entry_where)
            item = {key: check_text(entry.get(key, ""), entry_where + "." + key, key == "heading")
                    for key in ("heading", "detail", "date", "location")}
            check_date(item["date"], entry_where + ".date")
            item["bullets"] = check_list(entry.get("bullets", []), entry_where + ".bullets")
            item["paragraphs"] = check_list(entry.get("paragraphs", []), entry_where + ".paragraphs")
            item["emphasis"] = check_list(entry.get("emphasis", []), entry_where + ".emphasis")
            content = " ".join([item["heading"], item["detail"], *item["bullets"], *item["paragraphs"]])
            if any(term not in content for term in item["emphasis"]):
                raise ValueError(f"{entry_where}.emphasis must contain exact substrings of displayed content.")
            if any(re.match(r"^[•●]\s*|^[-*]\s+|^\d+[.)]\s+", b) for b in item["bullets"]):
                raise ValueError(f"{entry_where}.bullets: omit manual bullet/number prefixes.")
            clean["entries"].append(item)
        result["sections"].append(clean)
    if not any(s["kind"] in {"education", "experience", "projects", "campus"} for s in result["sections"]):
        raise ValueError("Include at least one education or experience/project/campus section.")
    return result


@dataclass(frozen=True)
class Block:
    kind: str
    text: str
    keep_next: bool = False
    columns: tuple[str, ...] = ()
    emphasis: tuple[str, ...] = ()
    column_fractions: tuple[float, ...] = ()


def make_blocks(profile: dict, layout: str = "reference-blue") -> list[Block]:
    blocks = [Block("name", profile["name"], True)]
    if profile["target_role"]:
        prefix = "求职意向：" if profile["language"] == "zh" else "Target role: "
        blocks.append(Block("role", prefix + profile["target_role"], True))
    if profile["contacts"]:
        contacts = profile["contacts"]
        current = []
        for contact in contacts:
            if current and (len(" | ".join(current + [contact])) > 65 or len(current) >= 2):
                blocks.append(Block("contact", " | ".join(current)))
                current = []
            current.append(contact)
        if current:
            blocks.append(Block("contact", " | ".join(current)))
    blocks.extend(Block("body", text) for text in profile["summary"])
    for section in profile["sections"]:
        blocks.append(Block("section", section["title"], True))
        for entry in section["entries"]:
            emphasis = tuple(entry.get("emphasis", []))
            paragraphs = entry.get("paragraphs", [])
            title = " | ".join(v for v in (entry["heading"], entry["detail"]) if v)
            meta = " | ".join(v for v in (entry["date"], entry["location"]) if v)
            body = bool(paragraphs or entry["bullets"])
            if layout == "reference-blue" and section["kind"] == "skills":
                label = entry["heading"] + ("：" if profile["language"] == "zh" else ": ")
                blocks.append(Block("bullet", label + entry["detail"], body, emphasis=(entry["heading"], *emphasis)))
            elif layout == "reference-blue" and entry["date"]:
                right = " · ".join(v for v in (entry["detail"], entry["location"]) if v)
                if section["kind"] == "education":
                    columns = (entry["heading"], right, entry["date"])
                    fractions = (0.24, 0.42, 0.34)
                else:
                    columns = (entry["date"], entry["heading"], right)
                    fractions = tuple(REFERENCE_STYLE["row_width_fractions"])
                blocks.append(Block("row", "\t".join(columns), body, columns, emphasis, fractions))
            else:
                blocks.append(Block("entry", title, bool(meta or body), emphasis=emphasis))
                if meta:
                    blocks.append(Block("meta", meta, body))
            blocks.extend(Block("narrative", text, emphasis=emphasis) for text in paragraphs)
            blocks.extend(Block("bullet", text, emphasis=emphasis) for text in entry["bullets"])
    return blocks


def content_review(profile: dict) -> dict:
    sparse = []
    for section in profile["sections"]:
        if section["kind"] not in {"experience", "projects"}:
            continue
        for entry in section["entries"]:
            texts = entry["bullets"] + entry.get("paragraphs", [])
            count = len("".join(texts).replace(" ", ""))
            if count < 60 or (len(texts) < 2 and count < 90):
                sparse.append({"section": section["title"], "entry": entry["heading"],
                               "body_characters": count, "reason": "Core experience needs concrete responsibilities, methods and delivery evidence."})
    return {"sparse_core_entries": sparse, "requires_followup": bool(sparse),
            "facts_checked": False, "note": "Length is a triage signal, not proof of writing quality."}


def text_segments(text: str, emphasis: tuple[str, ...]):
    if not emphasis:
        return [(text, False)]
    pattern = re.compile("|".join(re.escape(t) for t in sorted(set(emphasis), key=len, reverse=True)))
    segments = []
    cursor = 0
    for match in pattern.finditer(text):
        segments.extend([(text[cursor:match.start()], False), (match.group(), True)])
        cursor = match.end()
    segments.append((text[cursor:], False))
    return segments


def fits_row(columns: tuple[str, ...], size: float, width_pt: float, fractions: tuple[float, ...] = ()) -> bool:
    fractions = fractions or REFERENCE_STYLE["row_width_fractions"]
    def approximate_width(text):
        return sum(1 if ord(c) > 255 else 0.62 for c in text) * size
    return all(approximate_width(text) <= width_pt * fraction - 10
               for text, fraction in zip(columns, fractions))


def font_settings(style, font_name: str, size: float, bold: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    props = style.element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.append(fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), font_name)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)


def add_hyperlinked_text(paragraph, text: str, emphasis: tuple[str, ...] = (), accent: str = "000000") -> None:
    def add_plain(value):
        for segment, strong in text_segments(value, emphasis):
            run = paragraph.add_run(segment)
            if strong:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(accent)
    cursor = 0
    for match in URL.finditer(text):
        add_plain(text[cursor:match.start()])
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(match.group(), RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        props.append(color)
        run.append(props)
        node = OxmlElement("w:t")
        node.text = match.group()
        run.append(node)
        link.append(run)
        paragraph._p.append(link)
        cursor = match.end()
    add_plain(text[cursor:])


def bullet_numbering(document: Document) -> int:
    root = document.part.numbering_part.element
    abstract_id = max((int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))), default=0) + 1
    num_id = max((int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))), default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, val in (("start", "1"), ("numFmt", "bullet"), ("lvlText", "•"), ("lvlJc", "left")):
        element = OxmlElement("w:" + tag)
        element.set(qn("w:val"), val)
        level.append(element)
    abstract.append(level)
    root.append(abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), str(abstract_id))
    number.append(reference)
    root.append(number)
    return num_id


def paragraph_border(paragraph, color: str, width: int, side: str = "bottom", space: int = 1):
    props = paragraph._p.get_or_add_pPr()
    borders = props.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        props.append(borders)
    border = OxmlElement("w:" + side)
    for key, value in (("val", "single"), ("sz", str(width)), ("space", str(space)), ("color", color)):
        border.set(qn("w:" + key), value)
    borders.append(border)


def build_docx(blocks: list[Block], path: Path, font_name: str, size: float, layout: str = "reference-blue") -> None:
    doc = Document()
    ref = layout == "reference-blue"
    config = REFERENCE_STYLE
    accent = config["accent_hex"] if ref else "000000"
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = Mm(config["top_margin_mm"] if ref else 16)
    section.bottom_margin = Mm(config["bottom_margin_mm"] if ref else 16)
    section.left_margin = section.right_margin = Mm(config["side_margin_mm"] if ref else 16)
    for name, points, bold in (("Normal", size, False), ("Title", config["name_size_pt"] if ref else 20, True),
                               ("Heading 1", config["section_size_pt"] if ref else 12, True), ("Heading 2", size, True)):
        style = doc.styles[name]
        font_settings(style, font_name, points, bold)
        style.paragraph_format.line_spacing = config["body_line_multiplier"] if ref else 1.15
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(config["body_space_after_pt"] if ref else 3)
        style.font.color.rgb = RGBColor.from_string(accent if name in ("Title", "Heading 1") else config["body_hex"] if ref else "000000")
    if ref:
        section.header_distance = section.footer_distance = Mm(3)
        for style_name, area in (("Header", section.header), ("Footer", section.footer)):
            p = area.paragraphs[0]
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)
            font_settings(doc.styles[style_name], font_name, 1)
            paragraph_border(p, accent, 32, space=0)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = blocks[0].text + " Resume"
    doc.core_properties.subject = ""
    doc.core_properties.comments = ""
    num_id = bullet_numbering(doc)
    for block in blocks:
        style = {"name": "Title", "section": "Heading 1", "entry": "Heading 2", "row": "Heading 2"}.get(block.kind, "Normal")
        paragraph = doc.add_paragraph(style=style)
        if block.kind == "row":
            width_mm = 210 - 2 * (config["side_margin_mm"] if ref else 16)
            fractions = block.column_fractions or config["row_width_fractions"]
            if fits_row(block.columns, size, width_mm * mm - 12, fractions):
                paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(width_mm * fractions[0]), WD_TAB_ALIGNMENT.LEFT)
                paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(width_mm), WD_TAB_ALIGNMENT.RIGHT)
                display = "\t".join(block.columns)
            else:
                display = " | ".join(v for v in block.columns if v)
            add_hyperlinked_text(paragraph, display, block.emphasis, accent)
        else:
            add_hyperlinked_text(paragraph, block.text, block.emphasis, accent)
        fmt = paragraph.paragraph_format
        fmt.keep_with_next = block.keep_next
        fmt.widow_control = True
        if block.kind == "section":
            fmt.space_before, fmt.space_after = Pt(config["section_space_before_pt"] if ref else 8), Pt(config["section_space_after_pt"] if ref else 4)
            if ref:
                paragraph_border(paragraph, accent, 6)
        elif block.kind in ("entry", "row"):
            fmt.space_before = Pt(config["entry_space_before_pt"] if ref else 3)
        elif block.kind == "bullet":
            fmt.left_indent, fmt.first_line_indent = Mm(3.5), Mm(-2.8)
            props = paragraph._p.get_or_add_pPr()
            numbering = OxmlElement("w:numPr")
            for tag, val in (("ilvl", "0"), ("numId", str(num_id))):
                node = OxmlElement("w:" + tag)
                node.set(qn("w:val"), val)
                numbering.append(node)
            props.append(numbering)
    doc.save(path)


def font_candidates(language: str) -> list[Path]:
    windows = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    if language == "zh":
        values = [windows / "msyh.ttc", windows / "simhei.ttf",
                  Path("/System/Library/Fonts/STHeiti Light.ttc"),
                  Path("/System/Library/Fonts/PingFang.ttc"),
                  Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
                  Path("/usr/share/fonts/truetype/arphic/uming.ttc"),
                  Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")]
    else:
        values = [windows / "arial.ttf", Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
                  Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    return values


def load_pdf_font(blocks: list[Block], language: str, explicit: Path | None, index: int) -> tuple[str, str]:
    candidates = [explicit] if explicit else font_candidates(language)
    problems = []
    needed = set("".join(b.text for b in blocks) + "•")
    for path in candidates:
        if path is None or not path.is_file():
            continue
        try:
            font = TTFont("ResumeFont", str(path), subfontIndex=index)
            missing = sorted(c for c in needed if not c.isspace() and ord(c) not in font.face.charToGlyph)
            if missing:
                raise ValueError(f"missing glyphs: {''.join(missing[:16])}")
            pdfmetrics.registerFont(font)
            bold_name = "ResumeFont"
            bold_files = {
                "msyh.ttc": "msyhbd.ttc", "arial.ttf": "arialbd.ttf",
                "DejaVuSans.ttf": "DejaVuSans-Bold.ttf",
                "NotoSansCJK-Regular.ttc": "NotoSansCJK-Bold.ttc",
            }
            bold_file = path.with_name(bold_files.get(path.name, "__no_bold_font__"))
            if bold_file.is_file():
                try:
                    bold = TTFont("ResumeBold", str(bold_file), subfontIndex=index)
                    if all(c.isspace() or ord(c) in bold.face.charToGlyph for c in needed):
                        pdfmetrics.registerFont(bold)
                        bold_name = "ResumeBold"
                except Exception:
                    pass  # A usable regular face is sufficient; do not fake glyphs.
            pdfmetrics.registerFontFamily("ResumeFont", normal="ResumeFont", bold=bold_name, italic="ResumeFont", boldItalic=bold_name)
            return "ResumeFont", bold_name
        except Exception as exc:
            problems.append(f"{path}: {exc}")
    raise ValueError("No compatible embedded TrueType font covers the text. Supply --pdf-font and, for TTC, --font-index. " + " | ".join(problems))


def pdf_markup(text: str, emphasis: tuple[str, ...] = (), accent: str = "000000") -> str:
    def marked(value):
        return "".join(f'<b><font color="#{accent}">{escape(t)}</font></b>' if strong else escape(t)
                       for t, strong in text_segments(value, emphasis))
    parts = []
    cursor = 0
    for match in URL.finditer(text):
        parts.append(marked(text[cursor:match.start()]))
        value = escape(match.group(), {'"': "&quot;"})
        parts.append(f'<link href="{value}" color="black">{value}</link>')
        cursor = match.end()
    parts.append(marked(text[cursor:]))
    return "".join(parts)


class SectionParagraph(Paragraph):
    def draw(self):
        super().draw()
        self.canv.saveState()
        self.canv.setStrokeColor(HexColor("#" + REFERENCE_STYLE["accent_hex"]))
        self.canv.setLineWidth(REFERENCE_STYLE["section_rule_pt"])
        self.canv.line(0, -2, self.width, -2)
        self.canv.restoreState()


class EntryRow(Flowable):
    """A searchable three-field row with shared wrapping decisions in both outputs."""
    def __init__(self, block, style, size):
        super().__init__()
        self.block, self.style, self.size = block, style, size
        self.keepWithNext = block.keep_next
        self.spaceBefore, self.spaceAfter = style.spaceBefore, style.spaceAfter

    def wrap(self, available_width, available_height):
        self.width = available_width
        fractions = self.block.column_fractions or REFERENCE_STYLE["row_width_fractions"]
        if fits_row(self.block.columns, self.size, available_width, fractions):
            self.cells = []
            self.offsets = []
            offset = 0
            for i, (text, fraction) in enumerate(zip(self.block.columns, fractions)):
                style = ParagraphStyle(f"row_cell_{i}", parent=self.style, alignment=TA_RIGHT if i == 2 else TA_LEFT,
                                       spaceBefore=0, spaceAfter=0)
                paragraph = Paragraph(pdf_markup(text, self.block.emphasis, REFERENCE_STYLE["accent_hex"]), style)
                paragraph.wrap(available_width * fraction - (8 if i != 2 else 0), available_height)
                self.cells.append(paragraph)
                self.offsets.append(offset)
                offset += available_width * fraction
            self.height = max(p.height for p in self.cells)
        else:
            text = " | ".join(v for v in self.block.columns if v)
            self.cells = [Paragraph(pdf_markup(text, self.block.emphasis), self.style)]
            self.cells[0].wrap(available_width, available_height)
            self.offsets = [0]
            self.height = self.cells[0].height
        return self.width, self.height

    def draw(self):
        for paragraph, offset in zip(self.cells, self.offsets):
            paragraph.drawOn(self.canv, offset, self.height - paragraph.height)


def build_pdf(blocks: list[Block], path: Path, font: str, bold_font: str, size: float, language: str, layout: str = "reference-blue") -> None:
    ref = layout == "reference-blue"
    config = REFERENCE_STYLE
    accent = config["accent_hex"] if ref else "000000"
    base = dict(fontName=font, textColor=HexColor("#" + config["body_hex"]) if ref else "black", alignment=TA_LEFT,
                leading=size * (config["body_line_multiplier"] if ref else 1.15),
                spaceAfter=3, wordWrap="CJK" if language == "zh" else None,
                allowWidows=0, allowOrphans=0)
    styles = {}
    for kind in ("name", "role", "contact", "body", "narrative", "section", "entry", "row", "meta", "bullet"):
        settings = dict(base, fontSize=size)
        if kind in ("name", "section", "entry", "row"):
            settings["fontName"] = bold_font
        if kind == "name":
            settings.update(fontSize=config["name_size_pt"] if ref else 20, leading=27 if ref else 23, textColor=HexColor("#" + accent), spaceAfter=6)
        elif kind == "section":
            settings.update(fontSize=config["section_size_pt"] if ref else 12, leading=18 if ref else 13.8,
                            spaceBefore=config["section_space_before_pt"] if ref else 8,
                            spaceAfter=config["section_space_after_pt"] if ref else 4, textColor=HexColor("#" + accent))
        elif kind in ("entry", "row"):
            settings.update(spaceBefore=config["entry_space_before_pt"] if ref else 3)
        elif kind == "bullet":
            settings.update(leftIndent=3.5 * mm, bulletIndent=0.7 * mm, bulletFontName=font, bulletFontSize=size)
        styles[kind] = ParagraphStyle(kind, **settings)
    story = []
    for block in blocks:
        style = ParagraphStyle(block.kind + "_instance", parent=styles[block.kind], keepWithNext=block.keep_next)
        if block.kind == "row":
            story.append(EntryRow(block, style, size))
        else:
            cls = SectionParagraph if ref and block.kind == "section" else Paragraph
            story.append(cls(pdf_markup(block.text, block.emphasis, accent), style, bulletText="•" if block.kind == "bullet" else None))
    side_margin = config["side_margin_mm"] if ref else 16
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=side_margin * mm, rightMargin=side_margin * mm,
                            topMargin=(config["top_margin_mm"] if ref else 16) * mm,
                            bottomMargin=(config["bottom_margin_mm"] if ref else 16) * mm,
                            title=blocks[0].text + " Resume", author="")
    def page_frame(canvas, document):
        if not ref:
            return
        canvas.saveState()
        canvas.setStrokeColor(HexColor("#" + accent))
        canvas.setLineWidth(config["page_rule_width_pt"])
        left, right = side_margin * mm, A4[0] - side_margin * mm
        offset = config["page_rule_offset_mm"] * mm
        canvas.line(left, A4[1] - offset, right, A4[1] - offset)
        canvas.line(left, offset, right, offset)
        canvas.restoreState()
    doc.build(story, onFirstPage=page_frame, onLaterPages=page_frame)


def resolve_soffice(explicit: Path | None) -> str | None:
    if explicit:
        if not explicit.is_file():
            raise ValueError("--soffice must point to an existing executable.")
        return str(explicit.resolve())
    return shutil.which("soffice") or shutil.which("soffice.exe")


def convert_pdf(docx_path: Path, directory: Path, executable: str) -> Path:
    with tempfile.TemporaryDirectory(prefix="lo_profile_", dir=directory) as profile:
        command = [executable, "-env:UserInstallation=" + Path(profile).as_uri(),
                   "--headless", "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(directory), str(docx_path)]
        result = subprocess.run(command, capture_output=True, text=True, errors="replace", timeout=120)
        pdf = directory / (docx_path.stem + ".pdf")
        if result.returncode != 0 or not pdf.is_file() or pdf.stat().st_size == 0:
            raise RuntimeError("DOCX-to-PDF conversion failed: " + (result.stdout + result.stderr)[-2500:])
        return pdf


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--basename", default="resume")
    parser.add_argument("--pdf-engine", choices=("auto", "libreoffice", "reportlab"), default="auto")
    parser.add_argument("--layout", choices=("reference-blue", "ats"), default="reference-blue")
    parser.add_argument("--review-only", action="store_true", help="Print content gaps without generating files.")
    parser.add_argument("--allow-sparse", action="store_true", help="Export a user-requested fact-limited draft despite thin core experience.")
    parser.add_argument("--soffice", type=Path)
    parser.add_argument("--pdf-font", type=Path)
    parser.add_argument("--font-index", type=int, default=0)
    parser.add_argument("--font-name")
    parser.add_argument("--body-size", type=float, default=10.5)
    parser.add_argument("--max-pages", type=int, default=1)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    try:
        stem = args.basename
        if not stem or stem != stem.strip() or stem.endswith(".") or re.search(r'[<>:"/\\|?*\x00-\x1f]', stem):
            raise ValueError("--basename must be a filename stem without separators or reserved characters.")
        if re.fullmatch(r"(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?", stem, re.I):
            raise ValueError("--basename is reserved on Windows.")
        if not 10 <= args.body_size <= 12 or args.max_pages < 1 or args.font_index < 0:
            raise ValueError("Use body-size 10–12, max-pages >= 1 and font-index >= 0.")
        profile = validate_profile(json.loads(args.input.read_text(encoding="utf-8-sig")))
        review = content_review(profile)
        if args.review_only:
            print(json.dumps(review, ensure_ascii=False, indent=2))
            return 0
        if review["requires_followup"] and not args.allow_sparse:
            raise ValueError("Core experiences are too sparse for a polished resume. Ask for concrete responsibilities, methods and delivery evidence before export. "
                             "Use --review-only to inspect gaps; --allow-sparse is only for a user-requested fact-limited draft.")
        blocks = make_blocks(profile, args.layout)
        directory = args.output_dir.resolve()
        targets = [directory / (stem + extension) for extension in (".docx", ".pdf", ".qa.json")]
        if not args.overwrite and any(p.exists() for p in targets):
            raise FileExistsError("Output already exists; choose another basename or explicitly use --overwrite.")
        if args.input.resolve() in targets:
            raise ValueError("Input and output paths must differ.")
        executable = resolve_soffice(args.soffice) if args.pdf_engine != "reportlab" else None
        if args.pdf_engine == "libreoffice" and not executable:
            raise ValueError("LibreOffice is unavailable. Supply --soffice or explicitly choose reportlab.")
        engine = "libreoffice" if executable else "reportlab"
        font_name = args.font_name or ("Microsoft YaHei" if profile["language"] == "zh" else "Arial")
        warnings = []
        if review["requires_followup"]:
            warnings.append("Core experience remains sparse; this is a fact-limited draft, not a completed deep rewrite.")
        if engine == "reportlab":
            warnings.append("PDF is independently typeset from the same source; DOCX pagination is not verified and may differ.")
            pdf_font, pdf_bold_font = load_pdf_font(blocks, profile["language"], args.pdf_font, args.font_index)
        else:
            pdf_font = pdf_bold_font = ""
        directory.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".resume_build_", dir=directory) as temp:
            work = Path(temp)
            docx_path, pdf_path = work / (stem + ".docx"), work / (stem + ".pdf")
            build_docx(blocks, docx_path, font_name, args.body_size, args.layout)
            if engine == "libreoffice":
                convert_pdf(docx_path, work, executable)
            else:
                build_pdf(blocks, pdf_path, pdf_font, pdf_bold_font, args.body_size, profile["language"], args.layout)
            report = verify(docx_path, pdf_path)
            if not report["text_consistent"] or report["blank_pdf_page_indices"]:
                raise RuntimeError("Text verification failed; final outputs were not published. " + json.dumps(report))
            if report["pdf_pages"] > args.max_pages:
                warnings.append(f"PDF has {report['pdf_pages']} pages, exceeding target {args.max_pages}; edit content before reducing font size.")
            report.update(pdf_engine=engine, layout=args.layout, content_review=review,
                          docx_rendered_by_converter=engine == "libreoffice", warnings=warnings)
            # Recheck immediately before publishing to reduce accidental overwrite risk.
            if not args.overwrite and any(p.exists() for p in targets):
                raise FileExistsError("Output appeared while building; refusing to overwrite.")
            for source, target in zip((docx_path, pdf_path), targets[:2]):
                publish_output(source, target, overwrite=args.overwrite)
            final_check = verify(targets[0], targets[1])
            if not final_check["text_consistent"]:
                raise RuntimeError("Final output paths failed text verification.")
            report.update(final_check, published_files_reopened=True)
            if any(v is True for v in report["windows_acl_inheritance_disabled"].values()):
                raise RuntimeError("Final output has a protected Windows ACL; restore appropriate output-folder permissions before delivery.")
            qa_path = work / (stem + ".qa.json")
            qa_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            publish_output(qa_path, targets[2], overwrite=args.overwrite)
        for warning in warnings:
            print("WARNING: " + warning, file=sys.stderr)
        print(json.dumps({"docx": str(targets[0]), "pdf": str(targets[1]), "qa": str(targets[2]),
                          "pdf_pages": report["pdf_pages"], "pdf_engine": engine}, ensure_ascii=False, indent=2))
        return 0
    except Exception as exc:
        parser.exit(1, f"Resume export failed: {exc}\n")


if __name__ == "__main__":
    raise SystemExit(main())
